from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


FONT_NAME = "Arial"
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, size: float = 10, *, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT_NAME)
    rfonts.set(qn("w:hAnsi"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def insert_before(reference, paragraph) -> None:
    reference._p.addprevious(paragraph._p)


def add_cor_entry(input_path: Path, output_path: Path) -> None:
    doc = Document(input_path)
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    if any("COR Community Development Corporation" in p.text for p in paragraphs):
        raise ValueError("COR entry already exists")

    care_header = next(
        p
        for p in paragraphs
        if p.style.name == "Resume Job Header"
        and p.text.startswith("Saddleback College | CARE Corner Assistant")
    )
    bullet_template = next(p for p in paragraphs if p.style.name == "Resume Bullet")
    if bullet_template._p.pPr is None or bullet_template._p.pPr.numPr is None:
        raise ValueError("Existing bullet numbering definition was not found")

    header = doc.add_paragraph(style="Resume Job Header")
    run = header.add_run("COR Community Development Corporation")
    set_run_font(run, 10, bold=True)
    run = header.add_run(" | Civic Event Volunteer | Irvine, CA")
    set_run_font(run, 10)
    run = header.add_run("\tOct. 2025")
    set_run_font(run, 10)

    bullet_texts = [
        "Supported setup and attendee operations for the CA 2026 Gubernatorial Candidate Forum, including seating arrangements, registration preparation, and on-site logistics.",
        "Checked in attendees, provided directions, and helped maintain an organized guest flow throughout the public civic event.",
    ]
    new_paragraphs = [header]
    for text in bullet_texts:
        bullet = doc.add_paragraph(style="Resume Bullet")
        bullet._p.get_or_add_pPr().insert(0, deepcopy(bullet_template._p.pPr.numPr))
        run = bullet.add_run(text)
        set_run_font(run, 10)
        new_paragraphs.append(bullet)

    for paragraph in new_paragraphs:
        insert_before(care_header, paragraph)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    add_cor_entry(args.input, args.output)


if __name__ == "__main__":
    main()
