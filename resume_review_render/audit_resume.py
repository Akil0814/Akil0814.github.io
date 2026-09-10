import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


docx_path = Path(sys.argv[1])
doc = Document(docx_path)

job_headers = [p for p in doc.paragraphs if p.style.name == "Resume Job Header"]
bullets = [p for p in doc.paragraphs if p.style.name == "Resume Bullet"]
all_text = "\n".join(p.text for p in doc.paragraphs)

def has_numbering(paragraph):
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.numPr is not None

explicit_sizes = []
visible_fonts = set()
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.text:
            if run.font.size is not None:
                explicit_sizes.append(run.font.size.pt)
            if run.font.name:
                visible_fonts.add(run.font.name)

with zipfile.ZipFile(docx_path) as archive:
    names = set(archive.namelist())
    xml_parts = {
        name: archive.read(name).decode("utf-8", errors="replace")
        for name in names
        if name.endswith(".xml")
    }
    combined_xml = "\n".join(xml_parts.values())

result = {
    "path": str(docx_path),
    "sha256": hashlib.sha256(docx_path.read_bytes()).hexdigest().upper(),
    "tables": len(doc.tables),
    "headings": [p.text for p in doc.paragraphs if p.style.name == "Resume Section"],
    "job_header_count": len(job_headers),
    "job_header_tabs": [p.text.count("\t") for p in job_headers],
    "job_headers_numbered": [has_numbering(p) for p in job_headers],
    "bullet_count": len(bullets),
    "bullets_numbered": [has_numbering(p) for p in bullets],
    "minimum_explicit_font_pt": min(explicit_sizes) if explicit_sizes else None,
    "visible_fonts": sorted(visible_fonts),
    "replacement_character": "\ufffd" in all_text or "\ufffd" in combined_xml,
    "double_tab": "\t\t" in all_text,
    "fake_bullet_paragraphs": [p.text for p in doc.paragraphs if p.text.lstrip().startswith(("•", "·"))],
    "comment_parts": sorted(name for name in names if "comment" in name.lower()),
    "tracked_insertions": "<w:ins" in combined_xml,
    "tracked_deletions": "<w:del" in combined_xml,
    "hyperlink_targets": [],
    "cor_header_present": "COR Community Development Corporation | Civic Event Volunteer | Irvine, CA\tOct. 2025" in all_text,
    "cor_bullet_1_present": "Supported setup and attendee operations for the CA 2026 Gubernatorial Candidate Forum, including seating arrangements, registration preparation, and on-site logistics." in all_text,
    "cor_bullet_2_present": "Checked in attendees, provided directions, and helped maintain an organized guest flow throughout the public civic event." in all_text,
    "experience_order": [],
}

for rel in doc.part.rels.values():
    if rel.reltype.endswith("/hyperlink"):
        result["hyperlink_targets"].append(rel.target_ref)

experience_names = [
    "ANIME Impulse OC 2026",
    "CES 2026",
    "COR Community Development Corporation",
    "Saddleback College | CARE Corner Assistant",
    "UPS | Preload",
]
for name in experience_names:
    result["experience_order"].append(all_text.find(name))

print(json.dumps(result, indent=2, ensure_ascii=False))
